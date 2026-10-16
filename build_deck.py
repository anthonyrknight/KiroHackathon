"""Generate the Campus Form Assistant pitch deck as a .pptx file.

Run: py build_deck.py  ->  writes Campus_Form_Assistant.pptx
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE as SHAPE
from pptx.dml.color import RGBColor

# ---- palette (matches the app UI) ----------------------------------------
BG = RGBColor(0x0F, 0x17, 0x2A)        # slate-900
PANEL = RGBColor(0x1E, 0x29, 0x3B)     # slate-800
ACCENT = RGBColor(0x81, 0x8C, 0xF8)    # indigo-400
ACCENT_DK = RGBColor(0x63, 0x66, 0xF1)  # indigo-500
TEXT = RGBColor(0xE2, 0xE8, 0xF0)      # slate-200
MUTED = RGBColor(0x94, 0xA3, 0xB8)     # slate-400
OK = RGBColor(0x22, 0xC5, 0x5E)        # green-500
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SW, SH = Inches(13.333), Inches(7.5)   # 16:9

prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]


def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    bg.shadow.inherit = False
    # send to back
    s.shapes._spTree.remove(bg._element)
    s.shapes._spTree.insert(2, bg._element)
    return s


def accent_bar(s, top=Inches(0.0)):
    bar = s.shapes.add_shape(SHAPE.RECTANGLE, 0, top, Inches(0.18), SH)
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT_DK
    bar.line.fill.background()
    bar.shadow.inherit = False


def tb(s, left, top, width, height, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = s.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    return tf


def para(tf, text, size, color, bold=False, first=False, space_after=10,
         align=PP_ALIGN.LEFT, bullet=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = "Segoe UI"
    return p


def card(s, left, top, width, height, fill=PANEL):
    c = s.shapes.add_shape(SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    c.fill.solid()
    c.fill.fore_color.rgb = fill
    c.line.color.rgb = RGBColor(0x33, 0x41, 0x55)
    c.line.width = Pt(1)
    c.shadow.inherit = False
    try:
        c.adjustments[0] = 0.06
    except Exception:
        pass
    return c


# ==========================================================================
# Slide 1 — Title
# ==========================================================================
s = slide()
accent_bar(s)
tf = tb(s, Inches(1.0), Inches(2.2), Inches(11.3), Inches(2.5))
para(tf, "Campus Form Assistant", 54, WHITE, bold=True, first=True, space_after=6)
para(tf, "A conversational AI agent that replaces paper campus forms",
     26, ACCENT, space_after=4)
para(tf, "with guided workflows, rule validation, and ready-to-submit documents",
     22, MUTED)
tf2 = tb(s, Inches(1.0), Inches(6.3), Inches(11.3), Inches(0.8))
para(tf2, "Kiro Hackathon  ·  Group 5", 16, MUTED, first=True)

# ==========================================================================
# Slide 2 — The Problem
# ==========================================================================
s = slide()
accent_bar(s)
tf = tb(s, Inches(0.9), Inches(0.5), Inches(11.5), Inches(1.0))
para(tf, "The Problem", 36, WHITE, bold=True, first=True)
para(tf, "Campus paperwork is slow, error-prone, and frustrating", 20, ACCENT)

problems = [
    ("Paper & PDF forms", "Students download, print, hand-fill, and physically "
     "deliver forms. Fields are cramped and easy to misread."),
    ("No validation at the source", "Mistakes (wrong ID format, bad email, "
     "missing fields) aren't caught until a human reviews them later."),
    ("Hidden campus rules", "Deadlines, credit caps, and eligibility rules live "
     "in handbooks. Students don't know they broke a rule until they're rejected."),
    ("Slow back-and-forth", "A single error means the form bounces back, "
     "restarting the whole cycle and delaying registration."),
]
x = Inches(0.9)
y = Inches(1.9)
cw = Inches(5.6)
ch = Inches(2.2)
gap = Inches(0.25)
for i, (title, body) in enumerate(problems):
    col = i % 2
    row = i // 2
    left = x + col * (cw + gap)
    top = y + row * (ch + gap)
    card(s, left, top, cw, ch)
    t = tb(s, left + Inches(0.35), top + Inches(0.3), cw - Inches(0.7),
           ch - Inches(0.6))
    para(t, title, 20, ACCENT, bold=True, first=True, space_after=8)
    para(t, body, 15, TEXT, space_after=0)

# ==========================================================================
# Slide 3 — The Solution
# ==========================================================================
s = slide()
accent_bar(s)
tf = tb(s, Inches(0.9), Inches(0.5), Inches(11.5), Inches(1.0))
para(tf, "The Solution", 36, WHITE, bold=True, first=True)
para(tf, "Have a conversation instead of filling out a form", 20, OK)

steps = [
    ("1", "Guided conversation", "The agent asks one question at a time in plain "
     "language. The student answers naturally — 'add', 'yes', '11/15/2026'."),
    ("2", "Validate as you go", "Every answer is checked against campus rules "
     "immediately. Bad input is explained and re-asked on the spot."),
    ("3", "Enforce campus rules", "Credit caps, deadlines, ID and email formats, "
     "and cross-field rules are enforced before anything is submitted."),
    ("4", "Generate the document", "A submission-ready PDF is produced with a "
     "reference number, field table, and signature block."),
]
y = Inches(1.9)
for i, (num, title, body) in enumerate(steps):
    top = y + i * Inches(1.25)
    # number badge
    badge = s.shapes.add_shape(SHAPE.OVAL, Inches(0.9), top, Inches(0.8),
                               Inches(0.8))
    badge.fill.solid()
    badge.fill.fore_color.rgb = ACCENT_DK
    badge.line.fill.background()
    badge.shadow.inherit = False
    bt = badge.text_frame
    bt.word_wrap = False
    bp = bt.paragraphs[0]
    bp.alignment = PP_ALIGN.CENTER
    br = bp.add_run()
    br.text = num
    br.font.size = Pt(28)
    br.font.bold = True
    br.font.color.rgb = WHITE
    br.font.name = "Segoe UI"
    # text
    t = tb(s, Inches(2.0), top - Inches(0.05), Inches(10.3), Inches(1.2),
           anchor=MSO_ANCHOR.MIDDLE)
    para(t, title, 22, WHITE, bold=True, first=True, space_after=4)
    para(t, body, 16, MUTED, space_after=0)

# ==========================================================================
# Slide 4 — How it works (architecture)
# ==========================================================================
s = slide()
accent_bar(s)
tf = tb(s, Inches(0.9), Inches(0.5), Inches(11.5), Inches(1.0))
para(tf, "How It Works", 36, WHITE, bold=True, first=True)
para(tf, "A data-driven engine — new forms are config, not code", 20, ACCENT)

flow = [
    ("Chat UI", "Single-page web\nchat (no framework)"),
    ("Agent engine", "One-question flow\n+ validation"),
    ("Form schema", "Fields, validators\n& campus rules"),
    ("PDF document", "Submission-ready\noutput + reference"),
]
cw = Inches(2.7)
ch = Inches(1.8)
y = Inches(2.3)
start_x = Inches(0.9)
gap = Inches(0.45)
for i, (title, body) in enumerate(flow):
    left = start_x + i * (cw + gap)
    card(s, left, y, cw, ch, fill=PANEL)
    t = tb(s, left + Inches(0.2), y + Inches(0.25), cw - Inches(0.4),
           ch - Inches(0.5), anchor=MSO_ANCHOR.MIDDLE)
    para(t, title, 19, ACCENT, bold=True, first=True, space_after=8,
         align=PP_ALIGN.CENTER)
    para(t, body, 14, TEXT, align=PP_ALIGN.CENTER, space_after=0)
    if i < len(flow) - 1:
        arr = tb(s, left + cw, y, gap, ch, align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.MIDDLE)
        para(arr, "\u2192", 28, ACCENT, first=True, align=PP_ALIGN.CENTER)

tf = tb(s, Inches(0.9), Inches(4.7), Inches(11.5), Inches(2.3))
para(tf, "Key design choices", 20, WHITE, bold=True, first=True, space_after=10)
for line in [
    "Data-driven forms — add a new form (leave of absence, room booking) by "
    "defining one schema. The UI and PDF generator pick it up automatically.",
    "Deterministic core — validation and documents never depend on an LLM, so "
    "results are reliable and the app runs fully offline.",
    "Optional LLM polish — if an API key is set, the agent warms up its phrasing; "
    "if anything fails it falls back to built-in wording.",
]:
    para(tf, "•  " + line, 15, MUTED, space_after=8)

# ==========================================================================
# Slide 5 — Rules in action (concrete proof)
# ==========================================================================
s = slide()
accent_bar(s)
tf = tb(s, Inches(0.9), Inches(0.5), Inches(11.5), Inches(1.0))
para(tf, "Campus Rules In Action", 36, WHITE, bold=True, first=True)
para(tf, "Real validation from the working prototype", 20, ACCENT)

examples = [
    ("Credit cap", "Adding a course that pushes a student past the 21-credit "
     "limit is blocked — the agent bounces back and asks them to adjust."),
    ("Campus email only", "A personal gmail address is rejected; the agent "
     "re-asks for an address ending in @riverside.edu."),
    ("Format checks", "Student ID (S1234567) and course codes (CS 101) must "
     "match the expected pattern before moving on."),
    ("Late-change reason", "Drops after the deadline require a written reason "
     "for the registrar — enforced as a cross-field rule."),
]
x = Inches(0.9)
y = Inches(1.9)
cw = Inches(5.6)
ch = Inches(2.1)
gap = Inches(0.25)
for i, (title, body) in enumerate(examples):
    col = i % 2
    row = i // 2
    left = x + col * (cw + gap)
    top = y + row * (ch + gap)
    card(s, left, top, cw, ch)
    chk = tb(s, left + Inches(0.3), top + Inches(0.25), Inches(0.6), Inches(0.6))
    para(chk, "\u2713", 24, OK, bold=True, first=True)
    t = tb(s, left + Inches(0.95), top + Inches(0.3), cw - Inches(1.3),
           ch - Inches(0.6))
    para(t, title, 19, WHITE, bold=True, first=True, space_after=6)
    para(t, body, 14, MUTED, space_after=0)

# ==========================================================================
# Slide 6 — Proposed AWS Architecture (deployment plan)
# ==========================================================================
s = slide()
accent_bar(s)
tf = tb(s, Inches(0.9), Inches(0.5), Inches(11.5), Inches(1.0))
para(tf, "Proposed AWS Architecture", 36, WHITE, bold=True, first=True)
para(tf, "Deployment plan — mapping each component to an AWS service", 20, ACCENT)

aws_rows = [
    ("Amazon Bedrock", "Conversational phrasing",
     "Drop-in for the LLM layer already isolated in the agent — managed, "
     "no API keys to ship."),
    ("AWS App Runner / ECS Fargate", "Host the FastAPI app",
     "Run the existing web app as a container with autoscaling and HTTPS."),
    ("Amazon S3", "Store generated PDFs",
     "Persist submission-ready documents instead of holding them in memory."),
    ("Amazon DynamoDB / ElastiCache", "Session + workflow state",
     "Replace the in-memory session store so conversations survive restarts "
     "and scale out."),
    ("Amazon Cognito", "Student authentication",
     "Campus SSO / login so each submission is tied to a verified student."),
    ("Amazon API Gateway + CloudWatch", "Routing & observability",
     "Managed entry point for the chat API with logging, metrics, and alarms."),
]
x = Inches(0.9)
y = Inches(1.85)
cw = Inches(5.65)
ch = Inches(1.5)
gapx = Inches(0.2)
gapy = Inches(0.18)
for i, (svc, role, body) in enumerate(aws_rows):
    col = i % 2
    row = i // 2
    left = x + col * (cw + gapx)
    top = y + row * (ch + gapy)
    card(s, left, top, cw, ch)
    t = tb(s, left + Inches(0.3), top + Inches(0.2), cw - Inches(0.6),
           ch - Inches(0.4))
    p = t.paragraphs[0]
    p.space_after = Pt(4)
    r1 = p.add_run()
    r1.text = svc + "   "
    r1.font.size = Pt(17)
    r1.font.bold = True
    r1.font.color.rgb = ACCENT
    r1.font.name = "Segoe UI"
    r2 = p.add_run()
    r2.text = "· " + role
    r2.font.size = Pt(14)
    r2.font.italic = True
    r2.font.color.rgb = MUTED
    r2.font.name = "Segoe UI"
    para(t, body, 13, TEXT, space_after=0)

note = tb(s, Inches(0.9), Inches(6.95), Inches(11.5), Inches(0.45))
para(note, "Note: a proposed deployment architecture — the current prototype "
     "runs locally and is cloud-ready by design.", 12, MUTED, first=True,
     align=PP_ALIGN.CENTER)

# ==========================================================================
# Slide 7 — Impact / closing
# ==========================================================================
s = slide()
accent_bar(s)
tf = tb(s, Inches(0.9), Inches(0.6), Inches(11.5), Inches(1.0))
para(tf, "Why It Matters", 36, WHITE, bold=True, first=True)

metrics = [
    ("Fewer rejections", "Errors are caught at the source, not days later."),
    ("Faster for students", "A conversation beats printing and hand-filling."),
    ("Less staff load", "Clean, validated submissions reach the registrar."),
    ("Easy to extend", "Any paper form becomes a workflow in minutes."),
]
cw = Inches(2.75)
ch = Inches(2.4)
y = Inches(2.0)
start_x = Inches(0.9)
gap = Inches(0.3)
for i, (title, body) in enumerate(metrics):
    left = start_x + i * (cw + gap)
    card(s, left, y, cw, ch)
    t = tb(s, left + Inches(0.25), y + Inches(0.35), cw - Inches(0.5),
           ch - Inches(0.7), anchor=MSO_ANCHOR.TOP)
    para(t, title, 18, ACCENT, bold=True, first=True, space_after=10,
         align=PP_ALIGN.CENTER)
    para(t, body, 15, TEXT, align=PP_ALIGN.CENTER, space_after=0)

tf = tb(s, Inches(0.9), Inches(5.1), Inches(11.5), Inches(1.6),
        anchor=MSO_ANCHOR.MIDDLE)
para(tf, "Skip the paperwork. Just have a conversation.", 26, WHITE,
     bold=True, first=True, align=PP_ALIGN.CENTER, space_after=6)
para(tf, "Built and verified end-to-end: chat \u2192 validation \u2192 PDF download",
     16, MUTED, align=PP_ALIGN.CENTER)

prs.save("Campus_Form_Assistant.pptx")
print("Wrote Campus_Form_Assistant.pptx with", len(prs.slides._sldIdLst), "slides")
